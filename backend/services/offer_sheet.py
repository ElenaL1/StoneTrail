"""Parse a warehouse price sheet into promotion lines.

Column roles come from the header titles, not from their order.
A filled name without sizes is a group. Rows under it are sizes.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from typing import Any

OfferKind = str

FINISH_STEMS = (
    "полирован",
    "термо",
    "шлифован",
    "пилен",
    "колот",
    "бучард",
    "лощен",
    "лощён",
    "матов",
    "сатинирован",
)

KIND_STEMS = (
    ("плитк", "tile"),
    ("плит", "tile"),
    ("слэб", "slab"),
    ("слеб", "slab"),
    ("блок", "block"),
)


@dataclass
class ParsedLine:
    kind: OfferKind | None
    group_name: str
    stone_name: str
    label: str
    finish: str | None = None
    length_mm: int | None = None
    width_mm: int | None = None
    thickness_mm: int | None = None
    height_mm: int | None = None
    weight_kg: Decimal | None = None
    area_m2: Decimal | None = None
    price_amount: Decimal | None = None
    price_unit: str | None = None
    unresolved: bool = False
    issue: str | None = None


@dataclass
class ParsedSheet:
    lines: list[ParsedLine] = field(default_factory=list)
    offer_note: str = ""


@dataclass
class _Group:
    group_name: str
    stone_name: str
    finish: str | None
    thickness_mm: int | None


def parse_offer_sheet(data: bytes, filename: str) -> ParsedSheet:
    rows = _read_rows(data, filename)
    return parse_offer_rows(rows)


def parse_offer_rows(rows: list[list[Any]]) -> ParsedSheet:
    header_at, columns, price_unit = _find_header(rows)
    if header_at is None:
        return ParsedSheet(
            lines=[
                ParsedLine(
                    kind=None,
                    group_name="Прайс",
                    stone_name="",
                    label="Лист",
                    unresolved=True,
                    issue="В листе нет строки с колонками наименования и цены.",
                )
            ]
        )
    default_kind = _sheet_kind(set(columns.values()))
    lines: list[ParsedLine] = []
    notes: list[str] = []
    group: _Group | None = None
    started = False
    for raw in rows[header_at + 1 :]:
        cells = _row_map(raw, columns)
        if _is_blank(cells) and not _loose_text(raw):
            continue
        if _is_group(cells):
            title = str(cells.get("name") or "").strip()
            stone_name, finish, thickness = parse_group_title(title)
            group = _Group(title, stone_name, finish, thickness)
            continue
        if not _has_measure(cells):
            text = _loose_text(raw)
            if started and text:
                notes.append(text)
            continue
        started = True
        lines.append(_line_from(cells, group, default_kind, price_unit))
    return ParsedSheet(lines=lines, offer_note=" ".join(notes).strip())


def parse_group_title(value: str) -> tuple[str, str | None, int | None]:
    text = value.strip()
    text = re.sub(r"^(гр\.|гранит)\s*", "", text, flags=re.IGNORECASE)
    thickness = None
    thick = re.search(r"(\d+)\s*мм", text, flags=re.IGNORECASE)
    if thick:
        thickness = int(thick.group(1))
        text = text[: thick.start()] + text[thick.end() :]
    finish = None
    lowered = text.lower().replace("ё", "е")
    for stem in FINISH_STEMS:
        if stem.replace("ё", "е") in lowered:
            finish = stem
            text = re.sub(stem + r"\w*", "", text, flags=re.IGNORECASE)
            lowered = text.lower().replace("ё", "е")
            break
    stone = re.sub(r"\s+", " ", text).strip(" -–,")
    return stone, finish, thickness


def normalize_stone_name(value: str) -> str:
    text = value.lower().replace("ё", "е")
    text = re.sub(r"^(гр\.|гранит)\s+", "", text)
    return re.sub(r"\s+", " ", text).strip()


def _read_rows(data: bytes, filename: str) -> list[list[Any]]:
    name = filename.lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        return [
            [sheet.cell_value(row, col) for col in range(sheet.ncols)]
            for row in range(sheet.nrows)
        ]
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        if worksheet is None:
            return []
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _find_header(
    rows: list[list[Any]],
) -> tuple[int | None, dict[int, str], str | None]:
    for index, row in enumerate(rows[:30]):
        columns: dict[int, str] = {}
        price_unit = None
        for col, value in enumerate(row):
            role, unit = _header_role(value)
            if role and role not in columns.values():
                columns[col] = role
            if unit:
                price_unit = unit
        if "name" in columns.values() and (
            "price" in columns.values() or "length" in columns.values()
        ):
            return index, columns, price_unit
    return None, {}, None


def _header_role(value: Any) -> tuple[str | None, str | None]:
    if value is None:
        return None, None
    text = str(value).lower().replace("ё", "е").replace("²", "2")
    text = text.replace(".", " ").replace("/", " ").replace(",", " ")
    compact = "".join(text.split())
    spaced = " ".join(text.split())
    unit = None
    if "цена" in spaced or spaced in {"руб", "стоимость"}:
        role = "price"
        if "м2" in compact or "кв" in compact:
            unit = "m2"
        elif "слэб" in compact or "слеб" in compact:
            unit = "slab"
        elif "тон" in compact:
            unit = "ton"
        elif "шт" in compact or "блок" in compact:
            unit = "piece"
        return role, unit
    if "наименован" in spaced or spaced in {"название", "камень"}:
        return "name", None
    if spaced.startswith("длина"):
        return "length", None
    if spaced.startswith("ширина"):
        return "width", None
    if spaced.startswith("толщина"):
        return "thickness", None
    if spaced.startswith("высота"):
        return "height", None
    if "фактур" in spaced or "обработ" in spaced:
        return "finish", None
    if "площад" in spaced or ("кв" in compact and "м" in compact) or compact == "м2":
        return "area", None
    if spaced.startswith("вес"):
        return "weight", None
    if spaced in {"тип", "категория"}:
        return "kind", None
    return None, None


def _sheet_kind(columns: set[str]) -> OfferKind | None:
    if "height" in columns and "weight" in columns and "thickness" not in columns:
        return "block"
    if "area" in columns and "thickness" in columns:
        return "tile"
    if "thickness" in columns:
        return "slab"
    if "height" in columns:
        return "block"
    return None


def _row_map(row: list[Any], columns: dict[int, str]) -> dict[str, Any]:
    mapped: dict[str, Any] = {}
    for index, role in columns.items():
        if index < len(row):
            mapped[role] = row[index]
    return mapped


def _is_blank(cells: dict[str, Any]) -> bool:
    return all(_empty(value) for value in cells.values())


def _empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _is_group(cells: dict[str, Any]) -> bool:
    if _empty(cells.get("name")) or _has_measure(cells):
        return False
    title = str(cells.get("name") or "").strip()
    if not _looks_like_group(title) and (len(title) >= 40 or title.endswith(".")):
        return False
    return True


def _has_measure(cells: dict[str, Any]) -> bool:
    return any(
        not _empty(cells.get(key))
        for key in ("length", "width", "height", "price", "area", "weight")
    )


def _loose_text(row: list[Any]) -> str:
    parts = [
        str(value).strip() for value in row if isinstance(value, str) and value.strip()
    ]
    text = " ".join(parts)
    if len(text) < 12:
        return ""
    return text


def _line_from(
    cells: dict[str, Any],
    group: _Group | None,
    default_kind: OfferKind | None,
    price_unit: str | None,
) -> ParsedLine:
    explicit = _kind_cell(cells.get("kind"))
    kind = explicit or default_kind
    name = str(cells.get("name") or "").strip()
    length_mm, length_raw = _parse_mm(cells.get("length"))
    width_mm, width_raw = _parse_mm(cells.get("width"))
    thickness_mm, _ = _parse_mm(cells.get("thickness"))
    height_mm, _ = _parse_mm(cells.get("height"))
    finish = _text(cells.get("finish"))
    stone_name = ""
    group_name = group.group_name if group else name or "Позиция"
    if group and (
        not name
        or normalize_stone_name(name) == normalize_stone_name(group.group_name)
        or normalize_stone_name(name) == normalize_stone_name(group.stone_name)
    ):
        stone_name = group.stone_name
        finish = finish or group.finish
        thickness_mm = thickness_mm or group.thickness_mm
    elif group and name and not _looks_like_group(name):
        stone_name = group.stone_name
        finish = finish or group.finish
        thickness_mm = thickness_mm or group.thickness_mm
    else:
        parsed_stone, parsed_finish, parsed_thickness = parse_group_title(
            name or group_name
        )
        stone_name = parsed_stone
        finish = finish or parsed_finish
        thickness_mm = thickness_mm or parsed_thickness
        if not group:
            group_name = name or group_name
    label = _size_label(
        name, group, length_mm, length_raw, width_mm, width_raw, height_mm
    )
    unit = price_unit or _default_unit(kind)
    line = ParsedLine(
        kind=kind,
        group_name=group_name,
        stone_name=stone_name,
        label=label,
        finish=finish,
        length_mm=length_mm,
        width_mm=width_mm,
        thickness_mm=thickness_mm,
        height_mm=height_mm,
        weight_kg=_parse_decimal(cells.get("weight")),
        area_m2=_parse_decimal(cells.get("area")),
        price_amount=_parse_decimal(cells.get("price")),
        price_unit=unit,
    )
    _mark_issue(line)
    return line


def _looks_like_group(name: str) -> bool:
    lowered = name.lower().replace("ё", "е")
    if lowered.startswith("гр.") or lowered.startswith("гранит"):
        return True
    if re.search(r"\d+\s*мм", lowered):
        return True
    return any(stem.replace("ё", "е") in lowered for stem in FINISH_STEMS)


def _size_label(
    name: str,
    group: _Group | None,
    length_mm: int | None,
    length_raw: str | None,
    width_mm: int | None,
    width_raw: str | None,
    height_mm: int | None,
) -> str:
    own_name = name.strip()
    group_name = group.group_name if group else ""
    if own_name and normalize_stone_name(own_name) not in {
        normalize_stone_name(group_name),
        normalize_stone_name(group.stone_name) if group else "",
    }:
        if not _looks_like_group(own_name):
            return own_name
    parts: list[str] = []
    if length_raw:
        parts.append(length_raw)
    elif length_mm:
        parts.append(str(length_mm))
    if width_raw:
        parts.append(width_raw)
    elif width_mm:
        parts.append(str(width_mm))
    if height_mm and not length_raw:
        parts.append(str(height_mm))
    if parts:
        if length_raw:
            return " ".join(parts)
        return "×".join(parts)
    return own_name or "Позиция"


def _mark_issue(line: ParsedLine) -> None:
    if line.kind is None:
        line.unresolved = True
        line.issue = "Укажите вид: плита, слэб или блок."
        line.kind = "tile"
        return
    sized = any(
        value is not None for value in (line.length_mm, line.width_mm, line.height_mm)
    )
    if line.kind == "block" and not sized:
        line.unresolved = True
        line.issue = "У блока нет размеров."
        return
    if line.kind in {"tile", "slab"} and not sized and line.label in {"", "Позиция"}:
        line.unresolved = True
        line.issue = "Нет размеров."


def _kind_cell(value: Any) -> OfferKind | None:
    if _empty(value):
        return None
    text = str(value).lower().replace("ё", "е")
    for stem, kind in KIND_STEMS:
        if stem in text:
            return kind
    return None


def _default_unit(kind: OfferKind | None) -> str | None:
    if kind == "tile":
        return "m2"
    if kind == "slab":
        return "m2"
    if kind == "block":
        return "ton"
    return None


def _text(value: Any) -> str | None:
    if _empty(value):
        return None
    return str(value).strip()


def _parse_mm(value: Any) -> tuple[int | None, str | None]:
    if _empty(value):
        return None, None
    if isinstance(value, bool):
        return None, None
    if isinstance(value, (int, float, Decimal)):
        number = int(Decimal(str(value)))
        if number <= 0:
            return None, None
        return number, None
    text = str(value).strip()
    compact = text.replace(" ", "")
    if re.fullmatch(r"\d+", compact):
        return int(compact), None
    return None, text


def _parse_decimal(value: Any) -> Decimal | None:
    if _empty(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    text = re.sub(r"[^0-9.]", "", text)
    if not text or text == ".":
        return None
    return Decimal(text)
